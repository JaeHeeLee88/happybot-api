import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import pandas as pd
import math
from scipy.optimize import minimize
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

# ========================================
# BASE MODEL & CONSTANTS
# ========================================

BASE_CONDENSATION = 5.1    # kg/h
BASE_TUBES = 36           # EA
BASE_LENGTH = 580         # mm
BASE_CMM = 35.68          # CMM

# 단위 변환 상수 (1 mmAq = 9.80665 Pa)
MMAQ_TO_PA = 9.80665

FIN_CORRECTION = {
    "STS304": 1.00,
    "알루미늄": 1.18
}

# ========================================
# SciPy 수치해석 최적화 알고리즘
# ========================================

def optimize_condenser_scipy(
    material, tube_mat, fin_mat,
    moisture_in, feed_mass_day, discharge_mass_day, moisture_out, operating_hours_day,
    tube_od, tube_id, pitch_x, pitch_y, fin_pitch, max_dp_mmaq=15.0
):
    # 1. 1일 증발량 및 시간당 응축량 (kg/h)
    wi_day = feed_mass_day * (moisture_in / 100.0)
    wf_day = discharge_mass_day * (moisture_out / 100.0)
    evaporated_day = max(wi_day - wf_day, 0.0)
    condensation_rate = evaporated_day / operating_hours_day

    # 2. 핀 재질 및 형상 보정계수
    f_fin = FIN_CORRECTION.get(fin_mat, 1.0)
    base_load_ratio = condensation_rate / BASE_CONDENSATION
    target_load = base_load_ratio / f_fin

    od_factor = (12.7 / tube_od) ** 0.25
    pitch_factor = (pitch_x / 38.0) ** 0.12
    fin_dp_factor = (6.0 / fin_pitch) ** 0.5

    # mmAq 입력값을 Pa 단위로 변환 (계산 및 제약조건용)
    max_dp_pa = max_dp_mmaq * MMAQ_TO_PA

    # 3. SciPy 최적화 문제 정의
    # 목적함수: 체적(N_tubes * Length) 최소화
    def objective(x):
        n_tubes, length = x
        return n_tubes * length

    # 제약조건 1: 열용량 충족 (Capacity >= Target Load)
    def constraint_thermal(x):
        n_tubes, length = x
        capacity = ((n_tubes / BASE_TUBES) ** 0.95) * ((length / BASE_LENGTH) ** 0.60) / (od_factor * pitch_factor)
        return capacity - target_load

    # 제약조건 2: 압력손실 제한 (DP_Pa <= Max_DP_Pa)
    def constraint_dp(x):
        n_tubes, length = x
        dp_pa = 100.0 * (length / 580.0) ** 0.85 * (n_tubes / 36.0) ** 0.60 * fin_dp_factor
        return max_dp_pa - dp_pa

    # 초기 예측값 및 범위 설정
    x0 = [BASE_TUBES * (target_load ** 0.95) * od_factor, BASE_LENGTH * (target_load ** 0.60) * pitch_factor]
    bounds = [(10, 300), (200, 3000)]
    constraints = [
        {'type': 'ineq', 'fun': constraint_thermal},
        {'type': 'ineq', 'fun': constraint_dp}
    ]

    # SciPy SLSQP 최적화 수행
    res = minimize(objective, x0, method='SLSQP', bounds=bounds, constraints=constraints)

    if res.success:
        opt_tubes = math.ceil(res.x[0])
        opt_length = round(res.x[1])
    else:
        opt_tubes = math.ceil(x0[0])
        opt_length = round(x0[1])

    # 4. 최적 결과 기반 성능 지표 재계산
    cfd_factor = (
        (opt_length / BASE_LENGTH) ** (-0.08)
        * (opt_tubes / BASE_TUBES) ** (-0.04)
        * 0.95
    )
    efficiency = round(min(cfd_factor * 100.0, 100.0), 1)

    fan_cmm = round(BASE_CMM * (base_load_ratio ** 0.85), 1)
    
    # 계산된 압력손실 (mmAq 단위로 산출)
    dp_pa = 100.0 * (opt_length / 580.0) ** 0.85 * (opt_tubes / 36.0) ** 0.60 * fin_dp_factor
    dp_mmaq = round(dp_pa / MMAQ_TO_PA, 2)

    return {
        "원료명": material,
        "튜브재질": tube_mat,
        "핀재질": fin_mat,
        "원료함수율": moisture_in,
        "투입량": feed_mass_day,
        "배출량": discharge_mass_day,
        "배출함수율": moisture_out,
        "운전시간": operating_hours_day,
        "1일증발량(kg/day)": round(evaporated_day, 2),
        "시간당응축량(kg/h)": round(condensation_rate, 2),
        "보정부하율": round(target_load, 2),
        "튜브OD": tube_od,
        "튜브ID": tube_id,
        "PitchX": pitch_x,
        "PitchY": pitch_y,
        "FinPitch": fin_pitch,
        "허용압력손실(mmAq)": max_dp_mmaq,
        "튜브수(EA)": opt_tubes,
        "열교환기길이(mm)": opt_length,
        "핀보정계수": f_fin,
        "응축효율(%)": efficiency,
        "팬용량(CMM)": fan_cmm,
        "압력손실(mmAq)": dp_mmaq
    }


# ========================================
# GUI (사용자 인터페이스)
# ========================================

class CondenserApp:

    def __init__(self, root):
        self.root = root
        self.root.title("SciPy 수치해석 기반 응축열교환기 최적설계 (mmAq 단위 적용)")
        self.root.geometry("1180x720")

        self.results = []

        input_container = tk.Frame(root)
        input_container.pack(fill="x", padx=15, pady=10)

        # [좌측] 운전 및 공정 조건
        op_frame = tk.LabelFrame(input_container, text=" 공정 및 운전 조건 ", font=("맑은 고딕", 10, "bold"), padx=10, pady=10)
        op_frame.pack(side="left", fill="both", expand=True, padx=5)

        op_labels = [
            ("원료명", "슬러지_A"),
            ("원료함수율 (%)", "80"),
            ("투입량 (kg/day)", "10000"),
            ("배출량 (kg/day)", "3000"),
            ("배출함수율 (%)", "20"),
            ("일운전시간 (h/day)", "20")
        ]

        self.op_entries = {}
        for i, (label, default) in enumerate(op_labels):
            tk.Label(op_frame, text=label, anchor="w").grid(row=i, column=0, sticky="w", pady=2)
            entry = tk.Entry(op_frame, width=15)
            entry.insert(0, default)
            entry.grid(row=i, column=1, padx=5, pady=2)
            self.op_entries[label] = entry

        tk.Label(op_frame, text="튜브 재질", anchor="w").grid(row=6, column=0, sticky="w", pady=2)
        self.combo_tube_mat = ttk.Combobox(op_frame, values=["STS304", "알루미늄"], width=13, state="readonly")
        self.combo_tube_mat.set("STS304")
        self.combo_tube_mat.grid(row=6, column=1, padx=5, pady=2)

        tk.Label(op_frame, text="핀 재질", anchor="w").grid(row=7, column=0, sticky="w", pady=2)
        self.combo_fin_mat = ttk.Combobox(op_frame, values=["STS304", "알루미늄"], width=13, state="readonly")
        self.combo_fin_mat.set("알루미늄")
        self.combo_fin_mat.grid(row=7, column=1, padx=5, pady=2)

        # [우측] 튜브/Pitch 변수 및 mmAq 제약조건
        geo_frame = tk.LabelFrame(input_container, text=" 튜브/Pitch 변수 & SciPy 제약조건 ", font=("맑은 고딕", 10, "bold"), padx=10, pady=10)
        geo_frame.pack(side="right", fill="both", expand=True, padx=5)

        geo_labels = [
            ("튜브 OD (mm)", "12.7"),
            ("튜브 ID (mm)", "11.1"),
            ("Pitch X (mm)", "38.0"),
            ("Pitch Y (mm)", "32.0"),
            ("Fin Pitch (mm)", "6.0"),
            ("허용 압력손실 (mmAq)", "15.0")  # Pa -> mmAq 단위 수정 적용
        ]

        self.geo_entries = {}
        for i, (label, default) in enumerate(geo_labels):
            tk.Label(geo_frame, text=label, anchor="w").grid(row=i, column=0, sticky="w", pady=3)
            entry = tk.Entry(geo_frame, width=15)
            entry.insert(0, default)
            entry.grid(row=i, column=1, padx=5, pady=3)
            self.geo_entries[label] = entry

        # 버튼 영역
        btn_frame = tk.Frame(root)
        btn_frame.pack(pady=5)

        tk.Button(
            btn_frame,
            text="SciPy 최적화 계산 실행",
            command=self.add_result,
            width=22,
            height=2,
            bg="#1F4E78",
            fg="white",
            font=("맑은 고딕", 9, "bold")
        ).pack(side="left", padx=10)

        tk.Button(
            btn_frame,
            text="Excel 서식 저장",
            command=self.export_excel,
            width=15,
            height=2,
            bg="#2B579A",
            fg="white",
            font=("맑은 고딕", 9, "bold")
        ).pack(side="left", padx=10)

        # 결과 출력 (Treeview) - 열교환기 길이 및 mmAq 압력손실 명시
        columns = (
            "원료명", "튜브/핀 재질", "튜브OD/ID", "Pitch(X/Y)", "FinPitch", 
            "1일증발량", "시간당응축량", "최적 튜브수(EA)", "열교환기길이(mm)", 
            "응축효율(%)", "팬용량(CMM)", "압력손실(mmAq)"
        )

        self.tree = ttk.Treeview(root, columns=columns, show="headings", height=12)

        for col in columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=95, anchor="center")

        self.tree.column("열교환기길이(mm)", width=115, anchor="center")
        self.tree.column("최적 튜브수(EA)", width=110, anchor="center")
        self.tree.column("압력손실(mmAq)", width=110, anchor="center")
        self.tree.pack(padx=15, pady=10, fill="both", expand=True)

    def add_result(self):
        try:
            material = self.op_entries["원료명"].get().strip()
            if not material:
                messagebox.showwarning("입력 경고", "원료명을 입력해주세요.")
                return

            moisture_in = float(self.op_entries["원료함수율 (%)"].get())
            feed_mass_day = float(self.op_entries["투입량 (kg/day)"].get())
            discharge_mass_day = float(self.op_entries["배출량 (kg/day)"].get())
            moisture_out = float(self.op_entries["배출함수율 (%)"].get())
            operating_hours_day = float(self.op_entries["일운전시간 (h/day)"].get())

            tube_mat = self.combo_tube_mat.get()
            fin_mat = self.combo_fin_mat.get()

            tube_od = float(self.geo_entries["튜브 OD (mm)"].get())
            tube_id = float(self.geo_entries["튜브 ID (mm)"].get())
            pitch_x = float(self.geo_entries["Pitch X (mm)"].get())
            pitch_y = float(self.geo_entries["Pitch Y (mm)"].get())
            fin_pitch = float(self.geo_entries["Fin Pitch (mm)"].get())
            max_dp_mmaq = float(self.geo_entries["허용 압력손실 (mmAq)"].get())

            # SciPy 연동 최적화 수행
            res = optimize_condenser_scipy(
                material, tube_mat, fin_mat,
                moisture_in, feed_mass_day, discharge_mass_day, moisture_out, operating_hours_day,
                tube_od, tube_id, pitch_x, pitch_y, fin_pitch, max_dp_mmaq
            )

            self.results.append(res)

            self.tree.insert(
                "",
                "end",
                values=(
                    res["원료명"],
                    f"{res['튜브재질']}/{res['핀재질']}",
                    f"{res['튜브OD']}/{res['튜브ID']}",
                    f"{res['PitchX']}/{res['PitchY']}",
                    res["FinPitch"],
                    res["1일증발량(kg/day)"],
                    res["시간당응축량(kg/h)"],
                    res["튜브수(EA)"],
                    res["열교환기길이(mm)"],
                    res["응축효율(%)"],
                    res["팬용량(CMM)"],
                    res["압력손실(mmAq)"]
                )
            )

        except ValueError:
            messagebox.showerror("입력 오류", "수치 항목에는 올바른 숫자를 입력해주세요.")

    def export_excel(self):
        if not self.results:
            messagebox.showwarning("저장 경고", "저장할 데이터가 없습니다.")
            return

        filepath = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
            title="mmAq 단위 반영 최적화 Excel 저장"
        )

        if not filepath:
            return

        try:
            wb = openpyxl.Workbook()
            
            ws_config = wb.active
            ws_config.title = "기본_설정"
            
            base_config = [
                ["설명", "변수명", "설정값", "단위"],
                ["기준 시간당 응축량", "BASE_CONDENSATION", 5.1, "kg/h"],
                ["기준 튜브 수", "BASE_TUBES", 36, "EA"],
                ["기준 튜브 길이", "BASE_LENGTH", 580, "mm"],
                ["기준 팬 풍량", "BASE_CMM", 35.68, "CMM"],
                ["mmAq to Pa 변환상수", "MMAQ_TO_PA", 9.80665, "Pa/mmAq"]
            ]
            for r_idx, row in enumerate(base_config, start=1):
                for c_idx, val in enumerate(row, start=1):
                    ws_config.cell(row=r_idx, column=c_idx, value=val)

            material_db = [
                ["핀재질", "핀보정계수"],
                ["STS304", 1.00],
                ["알루미늄", 1.18]
            ]
            for r_idx, row in enumerate(material_db, start=1):
                for c_idx, val in enumerate(row, start=6):
                    ws_config.cell(row=r_idx, column=c_idx, value=val)

            ws_calc = wb.create_sheet(title="설계_계산")
            
            headers = [
                "원료명", "원료함수율(%)", "투입량(kg/day)", "배출량(kg/day)", "배출함수율(%)", 
                "일운전시간(h/day)", "튜브재질", "핀재질", "튜브OD(mm)", "튜브ID(mm)", 
                "PitchX(mm)", "PitchY(mm)", "FinPitch(mm)", "허용압력손실(mmAq)", "1일증발량(kg/day)", 
                "시간당응축량(kg/h)", "핀보정계수", "보정부하율", "최적 튜브수(EA)", "열교환기길이(mm)", 
                "CFD보정", "응축효율(%)", "팬용량(CMM)", "압력손실(mmAq)"
            ]
            ws_calc.append(headers)

            dv_mat = DataValidation(type="list", formula1='"STS304, 알루미늄"', allow_blank=True)
            ws_calc.add_data_validation(dv_mat)
            dv_mat.add(f"G2:H{len(self.results) + 100}")

            for idx, item in enumerate(self.results, start=2):
                row_data = [
                    item["원료명"],
                    item["원료함수율"],
                    item["투입량"],
                    item["배출량"],
                    item["배출함수율"],
                    item["운전시간"],
                    item["튜브재질"],
                    item["핀재질"],
                    item["튜브OD"],
                    item["튜브ID"],
                    item["PitchX"],
                    item["PitchY"],
                    item["FinPitch"],
                    item["허용압력손실(mmAq)"],
                    f"=MAX((C{idx}*(B{idx}/100)) - (D{idx}*(E{idx}/100)), 0)",                         # 1일증발량
                    f"=O{idx}/F{idx}",                                                                 # 시간당응축량
                    f"=VLOOKUP(H{idx}, 기본_설정!$F$2:$G$3, 2, FALSE)",                               # 핀보정계수
                    f"=(P{idx}/기본_설정!$C$2)/Q{idx}",                                                # 보정부하율
                    item["튜브수(EA)"],                                                                # SciPy 최적화값
                    item["열교환기길이(mm)"],                                                          # SciPy 최적화값
                    f"=((T{idx}/기본_설정!$C$4)^(-0.08))*((S{idx}/기본_설정!$C$3)^(-0.04))*0.95",         # CFD보정
                    f"=ROUND(MIN(U{idx}*100, 100), 1)",                                                # 응축효율(%)
                    f"=ROUND(기본_설정!$C$5*((P{idx}/기본_설정!$C$2)^0.85), 1)",                        # 팬용량(CMM)
                    f"=ROUND((100*(T{idx}/580)^0.85*(S{idx}/36)^0.6*(6/M{idx})^0.5)/9.80665, 2)"        # 압력손실(mmAq 변환)
                ]
                ws_calc.append(row_data)

            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            header_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
            align_center = Alignment(horizontal="center", vertical="center")

            for sheet in [ws_config, ws_calc]:
                for cell in sheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = align_center

                for col in sheet.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = openpyxl.utils.get_column_letter(col[0].column)
                    sheet.column_dimensions[col_letter].width = max(max_len + 3, 13)

            wb.save(filepath)
            messagebox.showinfo("성공", "mmAq 수치해석 결과 및 자동계산 서식이 포함된 엑셀 파일이 저장되었습니다.")

        except Exception as e:
            messagebox.showerror("오류", f"Excel 저장 중 오류 발생: {e}")


if __name__ == "__main__":
    root = tk.Tk()
    app = CondenserApp(root)
    root.mainloop()
